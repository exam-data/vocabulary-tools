#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx 释义换行：把表格中「释义」列按最大字符数断行（移植自旧仓 update_def/format_doc_def.py，
去掉数据库配置依赖，改为命令行参数；重写单元格时会套回生成器的格式，
不会像旧版那样洗掉美化——旧流程要跑两次 macro.js 就是因为这个）。

用法：
    python3 format_doc_def.py in.docx -o out.docx --max-length 20
    （max-length：释义单元格单行可容纳的最大汉字数；在 WPS 里实测后填入）
"""
import argparse, sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "华文中宋"

def str_add_break_line(zh_str, max_length):
    result = []
    zh_str = zh_str.strip()
    if not zh_str:
        return zh_str
    words = zh_str.split("、")
    current_length = 0
    for word in words:
        if len(word) > max_length:
            result.append(word)
            current_length = len(word) % max_length
            continue
        if current_length + len(word) + 1 <= max_length:
            if current_length > 0:
                result.append("、")
                current_length += 1
            result.append(word)
            current_length += len(word)
        else:
            if result:
                result.append("\n")
            result.append(word)
            current_length = len(word)
    if result and result[-1] == "、":
        result.pop()
    return "".join(result)

def rewrite_cell(cell, text):
    """重写单元格文本并套回格式（字体/居中/零间距），换行用真正的 <w:br/>"""
    cell.text = ""
    p = cell.paragraphs[0]
    for r in list(p.runs):                 # 清掉 cell.text 遗留的空 run
        r._element.getparent().remove(r._element)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    lines = text.split("\n")
    for i, seg in enumerate(lines):
        run = p.add_run(seg)
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if i < len(lines) - 1:
            run.add_break()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", help="输入 docx（export_docx.py 的产物）")
    ap.add_argument("-o", "--output", required=True, help="输出 docx")
    ap.add_argument("--max-length", type=int, required=True, help="释义单行最大汉字数")
    ap.add_argument("--column", default="释义", help="要处理的列标题（默认：释义）")
    args = ap.parse_args()

    doc = Document(args.input)
    hit = 0
    for table in doc.tables:
        header = {cell.text.strip(): i for i, cell in enumerate(table.rows[0].cells)}
        if args.column not in header:
            continue
        ci = header[args.column]
        for i, row in enumerate(table.rows):
            if i == 0:
                continue
            cell = row.cells[ci]
            rewrite_cell(cell, str_add_break_line(cell.text, args.max_length))
            hit += 1
    if not hit:
        sys.exit(f"没有找到「{args.column}」列，请检查表头")
    doc.save(args.output)
    print(f"ok: 处理了 {hit} 个单元格 -> {args.output}（格式已保留，无需再跑宏）")

if __name__ == "__main__":
    main()
